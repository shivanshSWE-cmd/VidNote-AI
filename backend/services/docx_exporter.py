import io
import datetime
import re
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_heading_styled(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)
    
    # Custom colors
    for run in heading.runs:
        if level == 1:
            run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Indigo #1E3A8A
            run.font.size = Pt(18)
            run.font.bold = True
        elif level == 2:
            run.font.color.rgb = RGBColor(14, 116, 144)  # Cyan #0E7490
            run.font.size = Pt(14)
            run.font.bold = True
    return heading

def parse_markdown_to_docx(doc, md_text: str):
    """Simple markdown parser converting headers, bold text, bullet points into docx paragraphs."""
    lines = md_text.split('\n')
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        if stripped.startswith('# '):
            add_heading_styled(doc, stripped[2:].strip(), level=1)
        elif stripped.startswith('## '):
            add_heading_styled(doc, stripped[3:].strip(), level=2)
        elif stripped.startswith('### '):
            add_heading_styled(doc, stripped[4:].strip(), level=3)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            add_formatted_text(p, stripped[2:].strip())
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            content = re.sub(r'^\d+\.\s', '', stripped)
            add_formatted_text(p, content)
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[2:].strip())
            run.font.italic = True
            run.font.color.rgb = RGBColor(100, 116, 139)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, stripped)

def add_formatted_text(paragraph, text: str):
    """Parses inline **bold** text and [timestamps]."""
    parts = re.split(r'(\*\*.*?\*\*|\[\d{2}:\d{2}\])', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('[') and part.endswith(']') and re.match(r'\[\d{2}:\d{2}\]', part):
            run = paragraph.add_run(f" {part} ")
            run.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235)  # Blue timestamp
        else:
            paragraph.add_run(part)

def create_docx_document(
    title: str,
    url: str,
    platform: str,
    duration_str: str,
    uploader: str,
    english_notes: str,
    hinglish_notes: str,
    transcript_text: str
) -> io.BytesIO:
    doc = Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Document Title Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run(title)
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("AI-Generated Video Notes & Transcript Summary")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)
    
    # Metadata Table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    metadata = [
        ("Source Platform", platform.capitalize()),
        ("Source URL", url),
        ("Creator / Channel", uploader),
        ("Video Duration", duration_str),
        ("Date Processed", datetime.datetime.now().strftime("%B %d, %Y - %H:%M"))
    ]
    
    for idx, (key, val) in enumerate(metadata):
        row = table.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        
        cell_k.width = Inches(2.0)
        cell_v.width = Inches(4.5)
        
        set_cell_background(cell_k, "F1F5F9")
        set_cell_background(cell_v, "F8FAFC")
        
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(key)
        rk.font.bold = True
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = RGBColor(51, 65, 85)
        
        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(val)
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = RGBColor(30, 41, 59)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Section 1: English Notes
    doc.add_heading("English Study Notes", level=1)
    parse_markdown_to_docx(doc, english_notes)
    
    doc.add_page_break()
    
    # Section 2: Hinglish Notes
    doc.add_heading("Hinglish Study Notes", level=1)
    parse_markdown_to_docx(doc, hinglish_notes)
    
    doc.add_page_break()
    
    # Section 3: Raw Transcript
    doc.add_heading("Verbatim Transcript with Timestamps", level=1)
    
    tp = doc.add_paragraph()
    tp.paragraph_format.space_after = Pt(6)
    trun = tp.add_run("Full chronological text transcript extracted from video:")
    trun.font.italic = True
    trun.font.color.rgb = RGBColor(71, 85, 105)
    
    t_lines = transcript_text.split('\n')
    for t_line in t_lines:
        if not t_line.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_formatted_text(p, t_line)
        
    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
